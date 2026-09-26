import os
import requests
import io
from fpdf import FPDF
from docx import Document
from docx.shared import Inches
import streamlit as st

BACKEND_URL = os.environ.get("BACKEND_URL", "http://backend:8000/generate")

def format_docx(text, document_type, raw_terms):
    doc = Document()
    try:
        doc.add_picture("Image/Logo.jpg", width=Inches(2))
    except Exception:
        pass 
        
    doc.add_heading(document_type, 0)

    for line in text.split("\n"):
        doc.add_paragraph(line)
        
    if raw_terms:
        doc.add_heading("Key Terms", level=2)
        table = doc.add_table(rows=0, cols=1)
        table.style = 'Table Grid' 
        
        term_list = raw_terms.split(";")
        for term in term_list:
            if term.strip():
                row_cells = table.add_row().cells
                row_cells[0].text = term.strip()
                
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

class BrandedPDF(FPDF):
    def header(self):
        try:
            self.image("Image/Logo.jpg", x=85, y=8, w=40)
        except Exception:
            pass
        self.ln(35)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", 'I', 8)
        self.cell(0, 10, "LegalEase INC. | Contact@legalease.com | All Rights Reserved", align='C')

def format_pdf(text, document_type):
    pdf = BrandedPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, txt=document_type, ln=True, align='C')

    pdf.set_font("Arial", size=12)
    clean_text = text.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 8, txt=clean_text)
    return pdf.output(dest='S').encode('latin-1')

st.set_page_config(page_title="LegalEase", layout="centered")

col1, col2, col3 = st.columns([1, 2, 1])
with col2:
    try:
        st.image("Image/Logo.jpg", use_container_width=True)
    except Exception:
        pass

st.markdown("<h2 style='text-align: center;'>AI Legal Document Generator</h2>", unsafe_allow_html=True)

document_type = st.text_input("Document Type (Ex: Agreement, Contract, NDA)")
parties = st.text_area("Parties Involved") 
terms = st.text_area("Terms and Conditions")
dates = st.text_input("Effective Dates")

if "generated_text" not in st.session_state:
    st.session_state.generated_text = ""    

if st.button("Generate Document"):
    # NEW: Added warning message for the free-tier server wake-up time
    st.info("💡 Note: The first document may take up to 60 seconds to generate while the secure server wakes up. Thanks for your patience!")
    
    with st.spinner("Drafting your legal document..."):                
        payload = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "dates": dates
        }
                
        try:
            # Uses the dynamic BACKEND_URL variable
            response = requests.post(BACKEND_URL, json=payload)
                
            if response.status_code == 200:                    
                st.session_state.generated_text = response.json().get("document", "")
                st.success("Document Generated Successfully!")
            else:
                st.error(f"Backend Error: {response.status_code}")
                    
        except requests.exceptions.ConnectionError:
            st.error("Could not connect to backend. Is FastAPI running?")

if st.session_state.generated_text:
    edited_text = st.text_area("Edit Document Below:", st.session_state.generated_text, height=300)
    
    btn1, btn2, btn3 = st.columns(3)
    with btn1:
        st.download_button(
            label="Download as .TXT",
            data=edited_text,
            file_name="legal_document.txt",
            mime="text/plain"
        )

    with btn2:
        st.download_button(
            label="Download as .DOCX",
            data=format_docx(edited_text, document_type, terms),
            file_name="legal_document.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    with btn3:
        st.download_button(
            label="Download as .PDF",
            data=format_pdf(edited_text, document_type),
            file_name="legal_document.pdf",
            mime="application/pdf"
        )